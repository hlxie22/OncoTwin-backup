#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import gzip
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Iterator


TEXT_SUFFIXES = (
    ".txt",
    ".tsv",
    ".csv",
    ".maf",
    ".txt.gz",
    ".tsv.gz",
    ".csv.gz",
    ".maf.gz",
)

PARQUET_SUFFIXES = (
    ".parquet",
    ".pq",
)

EXCEL_SUFFIXES = (
    ".xlsx",
    ".xlsm",
)


PATIENT_ID_EXACT = {
    "patient_id",
    "patientid",
    "patient",
    "pt_id",
    "ptid",
    "subject_id",
    "subjectid",
    "person_id",
    "record_id",
    "deid_patient_id",
    "deidentified_patient_id",
    "unique_patient_id",
}

SAMPLE_ID_EXACT = {
    "sample_id",
    "sampleid",
    "sample",
    "specimen_id",
    "specimenid",
    "tumor_sample_barcode",
    "tumor_sample_id",
    "aliquot_id",
    "unique_sample_id",
}

SITE_EXACT = {
    "center",
    "centre",
    "site",
    "institution",
    "hospital",
    "facility",
    "center_id",
    "site_id",
    "institution_id",
    "institution_name",
    "center_name",
}

CANCER_EXACT = {
    "cancer_type",
    "cancer_type_detailed",
    "primary_site",
    "tumor_type",
    "disease_type",
    "oncotree_code",
    "oncotree_primary_diagnosis_name",
    "primary_diagnosis",
    "diagnosis",
}

PANEL_EXACT = {
    "seq_assay_id",
    "assay_id",
    "panel_id",
    "gene_panel",
    "gene_panel_id",
    "sequencing_panel",
    "sequencing_assay",
    "assay",
    "panel",
}

DATE_TOKENS = (
    "date",
    "datetime",
    "timestamp",
    "time",
    "start",
    "stop",
    "end",
    "report",
    "result",
    "available",
    "availability",
    "collection",
    "collected",
    "diagnos",
    "death",
    "scan",
    "imaging",
    "specimen",
    "drawn",
)

SITE_MAP_PATTERNS = [
    (
        "MSK",
        re.compile(
            r"(^|[^A-Z])(?:MSK|MSKCC|MEMORIAL[ _-]*SLOAN[ _-]*KETTERING)([^A-Z]|$)",
            re.I,
        ),
    ),
    (
        "DFCI",
        re.compile(
            r"(^|[^A-Z])(?:DFCI|DANA[ _-]*FARBER)([^A-Z]|$)",
            re.I,
        ),
    ),
    (
        "VICC",
        re.compile(
            r"(^|[^A-Z])(?:VICC|VANDERBILT)([^A-Z]|$)",
            re.I,
        ),
    ),
]

BREAST_PAT = re.compile(
    r"(?:breast|\bbrca\b|mammary)",
    re.I,
)


@dataclass
class TableMeta:
    source: str
    path: Path
    relpath: str
    kind: str
    columns: list[str]
    normalized: dict[str, str]
    delimiter: str | None
    role: list[str]
    size_bytes: int
    header_error: str | None = None


def norm_col(name: str) -> str:
    return re.sub(
        r"[^a-z0-9]+",
        "_",
        str(name).strip().lower(),
    ).strip("_")


def file_kind(path: Path) -> str:
    name = path.name.lower()

    suffixes = sorted(
        TEXT_SUFFIXES + PARQUET_SUFFIXES + EXCEL_SUFFIXES,
        key=len,
        reverse=True,
    )

    for suffix in suffixes:
        if name.endswith(suffix):
            return suffix.lstrip(".")

    return "other"


def classify_role(path: Path) -> list[str]:
    s = str(path).lower()
    name = path.name.lower()

    roles: list[str] = []

    if (
        "clinical_patient" in s
        or re.search(r"(^|[_-])patient([_.-]|$)", name)
    ):
        roles.append("patient")

    if (
        "clinical_sample" in s
        or re.search(r"(^|[_-])sample([_.-]|$)", name)
    ):
        roles.append("sample")

    if any(
        token in s
        for token in (
            "treatment",
            "therapy",
            "regimen",
            "medication",
        )
    ):
        roles.append("treatment")

    if any(
        token in s
        for token in (
            "radiol",
            "imaging",
            "scan",
            "recist",
        )
    ):
        roles.append("radiology")

    if any(
        token in s
        for token in (
            "lab_test",
            "laboratory",
            "labs",
            "tumor_marker",
        )
    ):
        roles.append("lab")

    if "timeline" in s:
        roles.append("timeline")

    if any(
        token in s
        for token in (
            "survival",
            "outcome",
            "progression",
            "death",
            "pfs",
            "os_",
        )
    ):
        roles.append("outcome")

    if any(
        token in s
        for token in (
            "panel",
            "assay",
            "gene_panel",
        )
    ):
        roles.append("panel")

    if any(
        token in s
        for token in (
            "mutation",
            "mutations",
            "cna",
            "fusion",
            "sv_",
            "genomic",
        )
    ):
        roles.append("genomic")

    if name.startswith("meta_") or "/meta_" in s:
        roles.append("metadata")

    if not roles:
        roles.append("other")

    return sorted(set(roles))


def open_text(path: Path):
    if path.name.lower().endswith(".gz"):
        return gzip.open(
            path,
            "rt",
            encoding="utf-8-sig",
            errors="replace",
        )

    return path.open(
        "rt",
        encoding="utf-8-sig",
        errors="replace",
    )


def detect_delimiter(path: Path, line: str) -> str:
    name = path.name.lower()

    if name.endswith(
        (
            ".tsv",
            ".tsv.gz",
            ".maf",
            ".maf.gz",
        )
    ):
        return "\t"

    if name.endswith(
        (
            ".csv",
            ".csv.gz",
        )
    ):
        return ","

    counts = {
        "\t": line.count("\t"),
        ",": line.count(","),
        "|": line.count("|"),
    }

    if max(counts.values()) <= 0:
        return "\t"

    return max(
        counts,
        key=counts.get,
    )


def text_header(
    path: Path,
) -> tuple[list[str], str | None, str | None]:
    try:
        with open_text(path) as f:
            for line in f:
                if not line.strip():
                    continue

                if line.startswith("#"):
                    continue

                delimiter = detect_delimiter(
                    path,
                    line,
                )

                row = next(
                    csv.reader(
                        [line],
                        delimiter=delimiter,
                    )
                )

                columns = [
                    str(x).strip()
                    for x in row
                ]

                # cBioPortal metadata/config files are often key:value
                # documents rather than tables.
                if (
                    len(columns) == 1
                    and ":" in columns[0]
                    and "\t" not in line
                    and "," not in line
                ):
                    return [], None, None

                return columns, delimiter, None

        return [], None, None

    except Exception as exc:
        return (
            [],
            None,
            f"{type(exc).__name__}: {exc}",
        )


def parquet_header(
    path: Path,
) -> tuple[list[str], str | None]:
    try:
        import pyarrow.parquet as pq

        pf = pq.ParquetFile(path)

        return (
            list(pf.schema_arrow.names),
            None,
        )

    except Exception as exc:
        return (
            [],
            f"{type(exc).__name__}: {exc}",
        )


def excel_header(
    path: Path,
) -> tuple[list[str], str | None]:
    try:
        import openpyxl

        wb = openpyxl.load_workbook(
            path,
            read_only=True,
            data_only=True,
        )

        ws = wb[wb.sheetnames[0]]

        for row in ws.iter_rows(
            min_row=1,
            max_row=20,
            values_only=True,
        ):
            vals = [
                ""
                if x is None
                else str(x).strip()
                for x in row
            ]

            if any(vals):
                wb.close()
                return vals, None

        wb.close()

        return [], None

    except Exception as exc:
        return (
            [],
            f"{type(exc).__name__}: {exc}",
        )


def inspect_headers(
    source: str,
    root: Path,
) -> list[TableMeta]:
    out: list[TableMeta] = []

    for path in sorted(
        root.rglob("*")
    ):
        if (
            not path.is_file()
            or path.name.startswith(".")
        ):
            continue

        kind = file_kind(path)

        if kind == "other":
            continue

        columns: list[str] = []
        delimiter: str | None = None
        error: str | None = None

        if kind in {
            x.lstrip(".")
            for x in TEXT_SUFFIXES
        }:
            (
                columns,
                delimiter,
                error,
            ) = text_header(path)

        elif kind in {
            x.lstrip(".")
            for x in PARQUET_SUFFIXES
        }:
            (
                columns,
                error,
            ) = parquet_header(path)

        elif kind in {
            x.lstrip(".")
            for x in EXCEL_SUFFIXES
        }:
            (
                columns,
                error,
            ) = excel_header(path)

        normalized = {
            c: norm_col(c)
            for c in columns
        }

        out.append(
            TableMeta(
                source=source,
                path=path,
                relpath=str(
                    path.relative_to(root)
                ),
                kind=kind,
                columns=columns,
                normalized=normalized,
                delimiter=delimiter,
                role=classify_role(path),
                size_bytes=path.stat().st_size,
                header_error=error,
            )
        )

    return out


def exact_col(
    meta: TableMeta,
    names: set[str],
) -> str | None:
    for col, normalized in meta.normalized.items():
        if normalized in names:
            return col

    return None


def patient_col(
    meta: TableMeta,
) -> str | None:
    col = exact_col(
        meta,
        PATIENT_ID_EXACT,
    )

    if col:
        return col

    for col, normalized in meta.normalized.items():
        if (
            normalized.endswith("_patient_id")
            or normalized.startswith("patient_id_")
        ):
            return col

    return None


def sample_col(
    meta: TableMeta,
) -> str | None:
    col = exact_col(
        meta,
        SAMPLE_ID_EXACT,
    )

    if col:
        return col

    for col, normalized in meta.normalized.items():
        if (
            normalized.endswith("_sample_id")
            or normalized.endswith("_specimen_id")
        ):
            return col

    return None


def site_cols(
    meta: TableMeta,
) -> list[str]:
    cols: list[str] = []

    for col, normalized in meta.normalized.items():
        if (
            normalized in SITE_EXACT
            or any(
                token in normalized
                for token in (
                    "institution",
                    "center",
                    "centre",
                    "hospital",
                    "site_id",
                    "site_name",
                )
            )
        ):
            cols.append(col)

    return cols


def cancer_cols(
    meta: TableMeta,
) -> list[str]:
    cols: list[str] = []

    for col, normalized in meta.normalized.items():
        if (
            normalized in CANCER_EXACT
            or any(
                token in normalized
                for token in (
                    "cancer_type",
                    "tumor_type",
                    "primary_site",
                    "oncotree",
                )
            )
        ):
            cols.append(col)

    return cols


def panel_cols(
    meta: TableMeta,
) -> list[str]:
    cols: list[str] = []

    for col, normalized in meta.normalized.items():
        if (
            normalized in PANEL_EXACT
            or any(
                token in normalized
                for token in (
                    "seq_assay",
                    "panel_id",
                    "gene_panel",
                    "sequencing_panel",
                )
            )
        ):
            cols.append(col)

    return cols


def date_cols(
    meta: TableMeta,
) -> list[str]:
    out: list[str] = []

    non_temporal_suffixes = (
        "_type",
        "_status",
        "_state",
        "_class",
        "_code",
        "_name",
        "_flag",
        "_indicator",
        "_modality",
    )

    for col, normalized in meta.normalized.items():
        if normalized.endswith(
            non_temporal_suffixes
        ):
            continue

        if any(
            token in normalized
            for token in DATE_TOKENS
        ):
            out.append(col)

    return out


def cbioportal_column_descriptions(
    meta: TableMeta,
) -> dict[str, str]:
    """
    Best-effort extraction of cBioPortal comment-row descriptions.

    Typical clinical files contain:
      display names
      descriptions
      datatypes
      priorities
      actual header

    Only the file prefix is read. No outcome values are summarized.
    """

    if (
        meta.delimiter is None
        or meta.kind
        not in {
            x.lstrip(".")
            for x in TEXT_SUFFIXES
        }
    ):
        return {}

    comment_rows: list[list[str]] = []

    try:
        with open_text(meta.path) as f:
            for line in f:
                if not line.strip():
                    continue

                if line.startswith("#"):
                    payload = line[1:].rstrip(
                        "\r\n"
                    )

                    row = next(
                        csv.reader(
                            [payload],
                            delimiter=meta.delimiter,
                        )
                    )

                    comment_rows.append(row)
                    continue

                break

    except Exception:
        return {}

    matching = [
        row
        for row in comment_rows
        if len(row) == len(meta.columns)
    ]

    if not matching:
        return {}

    # Standard cBioPortal ordering is:
    # display name, description, datatype, priority.
    chosen = (
        matching[1]
        if len(matching) >= 2
        else matching[0]
    )

    return {
        meta.columns[i]: str(
            chosen[i]
        ).strip()
        for i in range(
            len(meta.columns)
        )
    }


def timestamp_semantic(
    col: str,
) -> str:
    normalized = norm_col(col)

    if any(
        token in normalized
        for token in (
            "report",
            "result",
            "available",
            "availability",
            "posted",
        )
    ):
        return "availability_time_candidate"

    if (
        "death" in normalized
        or "deceased" in normalized
    ):
        return "death_time"

    if "diagnos" in normalized:
        return "diagnosis_time"

    if any(
        token in normalized
        for token in (
            "collection",
            "collected",
            "specimen",
            "drawn",
        )
    ):
        return "specimen_or_draw_time"

    if any(
        token in normalized
        for token in (
            "scan",
            "imaging",
            "radiol",
        )
    ):
        return "scan_event_time"

    if any(
        token in normalized
        for token in (
            "start",
            "begin",
        )
    ):
        return "event_start_time"

    if any(
        token in normalized
        for token in (
            "stop",
            "end",
            "discontinu",
        )
    ):
        return "event_end_time"

    return "generic_event_time_candidate"


def should_full_scan(
    meta: TableMeta,
) -> bool:
    roles = set(meta.role)

    if not meta.columns:
        return False

    if meta.source == "CHORD":
        return bool(
            roles
            & {
                "patient",
                "sample",
                "treatment",
                "radiology",
                "lab",
                "timeline",
                "outcome",
            }
        )

    if meta.source == "BPC":
        # BPC is small enough to inspect all tabular files for IDs,
        # site linkage, and schema.
        #
        # We deliberately do NOT compute outcome-value distributions.
        return True

    if meta.source == "GENIE":
        # Avoid reading the full multi-GB mutation payload.
        # Clinical identity and panel files are sufficient for CKPT0D.
        return bool(
            roles
            & {
                "patient",
                "sample",
                "panel",
            }
        )

    return False


def iter_text_rows(
    meta: TableMeta,
) -> Iterator[dict[str, str]]:
    assert meta.delimiter is not None

    with open_text(meta.path) as f:

        def data_lines() -> Iterator[str]:
            for line in f:
                if not line.strip():
                    continue

                if line.startswith("#"):
                    continue

                yield line

        reader = csv.reader(
            data_lines(),
            delimiter=meta.delimiter,
        )

        try:
            header = [
                str(x).strip()
                for x in next(reader)
            ]
        except StopIteration:
            return

        for row in reader:
            if len(row) < len(header):
                row = row + [""] * (
                    len(header) - len(row)
                )

            elif len(row) > len(header):
                row = row[:len(header)]

            yield {
                header[i]: str(
                    row[i]
                ).strip()
                for i in range(
                    len(header)
                )
            }


def iter_parquet_rows(
    meta: TableMeta,
) -> Iterator[dict[str, str]]:
    import pyarrow.parquet as pq

    pf = pq.ParquetFile(
        meta.path
    )

    for batch in pf.iter_batches(
        batch_size=65536
    ):
        cols = batch.schema.names

        arrays = [
            batch.column(i).to_pylist()
            for i in range(
                len(cols)
            )
        ]

        for row_index in range(
            batch.num_rows
        ):
            yield {
                cols[col_index]:
                    ""
                    if arrays[col_index][row_index] is None
                    else str(
                        arrays[col_index][row_index]
                    ).strip()
                for col_index in range(
                    len(cols)
                )
            }


def iter_excel_rows(
    meta: TableMeta,
) -> Iterator[dict[str, str]]:
    import openpyxl

    wb = openpyxl.load_workbook(
        meta.path,
        read_only=True,
        data_only=True,
    )

    ws = wb[
        wb.sheetnames[0]
    ]

    header: list[str] | None = None

    for row in ws.iter_rows(
        values_only=True
    ):
        vals = [
            ""
            if x is None
            else str(x).strip()
            for x in row
        ]

        if header is None:
            if not any(vals):
                continue

            header = vals
            continue

        if not any(vals):
            continue

        if len(vals) < len(header):
            vals += [""] * (
                len(header) - len(vals)
            )

        yield {
            header[i]: vals[i]
            for i in range(
                len(header)
            )
        }

    wb.close()


def iter_rows(
    meta: TableMeta,
) -> Iterator[dict[str, str]]:
    if meta.kind in {
        x.lstrip(".")
        for x in TEXT_SUFFIXES
    }:
        yield from iter_text_rows(meta)

    elif meta.kind in {
        x.lstrip(".")
        for x in PARQUET_SUFFIXES
    }:
        yield from iter_parquet_rows(meta)

    elif meta.kind in {
        x.lstrip(".")
        for x in EXCEL_SUFFIXES
    }:
        yield from iter_excel_rows(meta)


def map_site(
    value: str,
) -> str | None:
    text = str(value).strip()

    if not text:
        return None

    for site, pattern in SITE_MAP_PATTERNS:
        if pattern.search(text):
            return site

    return None


def infer_site_from_patient_id(
    patient_id: str,
) -> str | None:
    return map_site(
        patient_id
    )


def is_breast_value(
    value: str,
) -> bool:
    return bool(
        BREAST_PAT.search(
            str(value)
        )
    )


def summarize_gene_panel_definition(
    meta: TableMeta,
) -> dict[str, Any] | None:
    if "panel" not in meta.role:
        return None

    result: dict[str, Any] = {
        "path": meta.relpath,
        "gene_count": None,
        "panel_id": None,
    }

    genes: set[str] = set()

    try:
        if meta.kind in {
            x.lstrip(".")
            for x in TEXT_SUFFIXES
        }:
            with open_text(meta.path) as f:
                for index, line in enumerate(f):
                    if index > 5000:
                        break

                    stripped = line.strip()
                    low = stripped.lower()

                    if (
                        low.startswith("stable_id:")
                        or low.startswith("gene_panel_id:")
                        or low.startswith("panel_id:")
                    ):
                        result["panel_id"] = (
                            stripped
                            .split(":", 1)[1]
                            .strip()
                        )

                    if low.startswith(
                        "gene_list:"
                    ):
                        payload = (
                            stripped
                            .split(":", 1)[1]
                        )

                        for token in re.split(
                            r"[\s,;]+",
                            payload,
                        ):
                            token = token.strip()

                            if token:
                                genes.add(token)

            if (
                not genes
                and meta.columns
            ):
                gene_columns = [
                    col
                    for col, normalized
                    in meta.normalized.items()
                    if normalized
                    in {
                        "hugo_symbol",
                        "gene",
                        "gene_symbol",
                    }
                ]

                if (
                    gene_columns
                    and meta.delimiter
                ):
                    gene_col = (
                        gene_columns[0]
                    )

                    for row in iter_rows(
                        meta
                    ):
                        value = (
                            row
                            .get(
                                gene_col,
                                "",
                            )
                            .strip()
                        )

                        if value:
                            genes.add(value)

        if genes:
            result["gene_count"] = (
                len(genes)
            )

        if not result["panel_id"]:
            stem = meta.path.name

            stem = re.sub(
                r"\.(txt|tsv|csv|gz)+$",
                "",
                stem,
                flags=re.I,
            )

            result["panel_id"] = stem

        return result

    except Exception as exc:
        result["error"] = (
            f"{type(exc).__name__}: "
            f"{exc}"
        )

        return result


def analyze_table(
    meta: TableMeta,
) -> tuple[
    dict[str, Any],
    dict[str, Any],
]:
    pcol = patient_col(meta)
    scol = sample_col(meta)

    scols = site_cols(meta)
    ccols = cancer_cols(meta)
    pcols = panel_cols(meta)
    dcols = date_cols(meta)

    patient_ids: set[str] = set()
    sample_ids: set[str] = set()

    patient_site_votes: dict[
        str,
        Counter[str],
    ] = defaultdict(Counter)

    breast_patients: set[str] = set()

    site_values: Counter[str] = Counter()
    cancer_values: Counter[str] = Counter()
    panel_values: Counter[str] = Counter()

    rows = 0
    parse_error: str | None = None

    try:
        for row in iter_rows(meta):
            rows += 1

            pid = (
                row.get(
                    pcol,
                    "",
                ).strip()
                if pcol
                else ""
            )

            sid = (
                row.get(
                    scol,
                    "",
                ).strip()
                if scol
                else ""
            )

            if pid:
                patient_ids.add(pid)

            if sid:
                sample_ids.add(sid)

            # BPC external-policy guard:
            # Only ID and site information is consumed here.
            # No outcome-value distributions are accumulated.
            if (
                meta.source == "BPC"
                and pid
            ):
                for col in scols:
                    raw = (
                        row
                        .get(
                            col,
                            "",
                        )
                        .strip()
                    )

                    mapped = map_site(raw)

                    if raw:
                        site_values[raw] += 1

                    if mapped:
                        patient_site_votes[
                            pid
                        ][mapped] += 1

            if (
                meta.source == "CHORD"
                and pid
            ):
                for col in ccols:
                    raw = (
                        row
                        .get(
                            col,
                            "",
                        )
                        .strip()
                    )

                    if raw:
                        cancer_values[
                            raw
                        ] += 1

                        if is_breast_value(
                            raw
                        ):
                            breast_patients.add(
                                pid
                            )

            if (
                meta.source == "GENIE"
                and (
                    scol
                    or "sample" in meta.role
                )
            ):
                for col in pcols:
                    raw = (
                        row
                        .get(
                            col,
                            "",
                        )
                        .strip()
                    )

                    if raw:
                        panel_values[
                            raw
                        ] += 1

    except Exception as exc:
        parse_error = (
            f"{type(exc).__name__}: "
            f"{exc}"
        )

    summary = {
        "source": meta.source,
        "path": meta.relpath,
        "roles": meta.role,
        "size_bytes": meta.size_bytes,
        "row_count": rows,
        "patient_id_column": pcol,
        "sample_id_column": scol,
        "unique_patients": len(
            patient_ids
        ),
        "unique_samples": len(
            sample_ids
        ),
        "site_columns": scols,
        "cancer_columns": ccols,
        "panel_columns": pcols,
        "date_columns": dcols,
        "parse_error": parse_error,
    }

    private = {
        "patient_ids": patient_ids,
        "sample_ids": sample_ids,
        "patient_site_votes":
            patient_site_votes,
        "breast_patients":
            breast_patients,
        "site_values": site_values,
        "cancer_values":
            cancer_values,
        "panel_values": panel_values,
    }

    return summary, private


def authoritative_score(
    meta: TableMeta,
    entity: str,
) -> tuple[int, int, str]:
    name = meta.path.name.lower()
    score = 0

    if entity == "patient":
        if (
            "data_clinical_patient"
            in name
        ):
            score += 100

        if "patient" in meta.role:
            score += 30

        if patient_col(meta):
            score += 20

    else:
        if (
            "data_clinical_sample"
            in name
        ):
            score += 100

        if "sample" in meta.role:
            score += 30

        if sample_col(meta):
            score += 20

        if patient_col(meta):
            score += 5

    return (
        -score,
        len(meta.relpath),
        meta.relpath,
    )


def choose_authoritative(
    metas: list[TableMeta],
    entity: str,
) -> list[TableMeta]:
    candidates: list[
        TableMeta
    ] = []

    for meta in metas:
        if (
            entity == "patient"
            and patient_col(meta)
            and "patient" in meta.role
        ):
            candidates.append(meta)

        elif (
            entity == "sample"
            and sample_col(meta)
            and "sample" in meta.role
        ):
            candidates.append(meta)

    if not candidates:
        for meta in metas:
            if (
                entity == "patient"
                and patient_col(meta)
            ):
                candidates.append(meta)

            elif (
                entity == "sample"
                and sample_col(meta)
            ):
                candidates.append(meta)

    candidates.sort(
        key=lambda meta:
            authoritative_score(
                meta,
                entity,
            )
    )

    return candidates


def pair_overlap(
    left: set[str],
    right: set[str],
) -> dict[str, Any]:
    intersection = (
        left
        & right
    )

    return {
        "left_count": len(left),
        "right_count": len(right),
        "exact_overlap_count":
            len(intersection),
        "left_overlap_fraction":
            (
                round(
                    len(intersection)
                    / len(left),
                    6,
                )
                if left
                else None
            ),
        "right_overlap_fraction":
            (
                round(
                    len(intersection)
                    / len(right),
                    6,
                )
                if right
                else None
            ),
    }


def write_tsv(
    path: Path,
    header: list[str],
    rows: Iterable[
        Iterable[Any]
    ],
) -> None:
    tmp = path.with_suffix(
        path.suffix + ".tmp"
    )

    with tmp.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as f:
        writer = csv.writer(
            f,
            delimiter="\t",
        )

        writer.writerow(header)

        for row in rows:
            writer.writerow(
                list(row)
            )

    if tmp.stat().st_size == 0:
        raise RuntimeError(
            f"empty TSV write: "
            f"{path}"
        )

    tmp.replace(path)


def atomic_json(
    path: Path,
    obj: Any,
) -> None:
    tmp = path.with_suffix(
        path.suffix + ".tmp"
    )

    tmp.write_text(
        json.dumps(
            obj,
            indent=2,
            sort_keys=True,
        ),
        encoding="utf-8",
    )

    if tmp.stat().st_size == 0:
        raise RuntimeError(
            f"empty JSON write: "
            f"{path}"
        )

    # Validate before atomic replacement.
    json.loads(
        tmp.read_text(
            encoding="utf-8"
        )
    )

    tmp.replace(path)


def atomic_text(
    path: Path,
    text: str,
) -> None:
    tmp = path.with_suffix(
        path.suffix + ".tmp"
    )

    tmp.write_text(
        text,
        encoding="utf-8",
    )

    if tmp.stat().st_size == 0:
        raise RuntimeError(
            f"empty text write: "
            f"{path}"
        )

    tmp.replace(path)


def main() -> int:
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--repo-root",
        default=".",
    )

    parser.add_argument(
        "--output-dir",
        default=(
            "artifacts/"
            "checkpoint0d"
        ),
    )

    args = parser.parse_args()

    repo = Path(
        args.repo_root
    ).resolve()

    output_arg = Path(
        args.output_dir
    )

    out = (
        output_arg
        if output_arg.is_absolute()
        else (
            repo
            / output_arg
        )
    ).resolve()

    out.mkdir(
        parents=True,
        exist_ok=True,
    )

    roots = {
        "GENIE":
            repo
            / "data"
            / "external_sources"
            / "genie_20_0_public",

        "CHORD":
            repo
            / "data"
            / "external_sources"
            / "msk_chord_2024"
            / "raw",

        "BPC":
            repo
            / "data"
            / "external_sources"
            / "bpc_brca_1_0_public",
    }

    missing = [
        f"{name}:{path}"
        for name, path
        in roots.items()
        if not path.exists()
    ]

    if missing:
        raise SystemExit(
            "[CKPT0D_FATAL] "
            "missing source roots: "
            + ", ".join(missing)
        )

    ###########################################################################
    # Phase A: schema/header inventory only.
    ###########################################################################

    print(
        "[CKPT0D] "
        "Header/schema inventory starting",
        flush=True,
    )

    all_metas: dict[
        str,
        list[TableMeta],
    ] = {}

    for source, root in roots.items():
        metas = inspect_headers(
            source,
            root,
        )

        all_metas[
            source
        ] = metas

        print(
            "[CKPT0D] "
            f"source={source} "
            f"files_inspected={len(metas)}",
            flush=True,
        )

    schema_path = (
        out
        / "schema_inventory.jsonl"
    )

    schema_tmp = (
        schema_path
        .with_suffix(
            ".jsonl.tmp"
        )
    )

    with schema_tmp.open(
        "w",
        encoding="utf-8",
    ) as f:
        for source in (
            "GENIE",
            "CHORD",
            "BPC",
        ):
            for meta in all_metas[
                source
            ]:
                row = {
                    "source":
                        meta.source,
                    "path":
                        meta.relpath,
                    "kind":
                        meta.kind,
                    "size_bytes":
                        meta.size_bytes,
                    "roles":
                        meta.role,
                    "columns":
                        meta.columns,
                    "patient_id_column":
                        patient_col(meta),
                    "sample_id_column":
                        sample_col(meta),
                    "site_columns":
                        site_cols(meta),
                    "cancer_columns":
                        cancer_cols(meta),
                    "panel_columns":
                        panel_cols(meta),
                    "date_columns":
                        date_cols(meta),
                    "column_descriptions":
                        cbioportal_column_descriptions(
                            meta
                        ),
                    "header_error":
                        meta.header_error,
                }

                f.write(
                    json.dumps(
                        row,
                        sort_keys=True,
                    )
                    + "\n"
                )

    if (
        not schema_tmp.exists()
        or schema_tmp.stat().st_size == 0
    ):
        raise RuntimeError(
            "schema inventory is empty"
        )

    schema_tmp.replace(
        schema_path
    )

    ###########################################################################
    # Phase B: choose authoritative identity files from schema evidence.
    ###########################################################################

    authority: dict[
        str,
        Any,
    ] = {}

    authority_metas: dict[
        str,
        dict[
            str,
            list[TableMeta],
        ],
    ] = {}

    for source in (
        "GENIE",
        "CHORD",
        "BPC",
    ):
        patient_files = (
            choose_authoritative(
                all_metas[source],
                "patient",
            )
        )

        sample_files = (
            choose_authoritative(
                all_metas[source],
                "sample",
            )
        )

        authority_metas[
            source
        ] = {
            "patient":
                patient_files,
            "sample":
                sample_files,
        }

        authority[
            source
        ] = {
            "patient_files": [
                x.relpath
                for x
                in patient_files
            ],
            "sample_files": [
                x.relpath
                for x
                in sample_files
            ],
            "patient_id_fields":
                sorted(
                    {
                        patient_col(x)
                        for x
                        in patient_files
                        if patient_col(x)
                    }
                ),
            "sample_id_fields":
                sorted(
                    {
                        sample_col(x)
                        for x
                        in sample_files
                        if sample_col(x)
                    }
                ),
            "selection_rule": (
                "standard clinical patient/sample "
                "filename + exact ID-field ranking; "
                "unions are used when a release is "
                "split across centers"
            ),
        }

    atomic_json(
        out
        / "authoritative_entities.json",
        authority,
    )

    ###########################################################################
    # Phase C: full scans only of key tables.
    ###########################################################################

    print(
        "[CKPT0D] "
        "Full scans of selected "
        "key tables starting",
        flush=True,
    )

    full_summaries: list[
        dict[str, Any]
    ] = []

    private_by_key: dict[
        tuple[str, str],
        dict[str, Any],
    ] = {}

    for source in (
        "GENIE",
        "CHORD",
        "BPC",
    ):
        must_scan_paths = {
            meta.relpath
            for entity in (
                "patient",
                "sample",
            )
            for meta
            in authority_metas[
                source
            ][entity]
        }

        selected = [
            meta
            for meta
            in all_metas[source]
            if (
                should_full_scan(meta)
                or meta.relpath
                in must_scan_paths
            )
        ]

        print(
            "[CKPT0D] "
            f"source={source} "
            f"full_scan_tables="
            f"{len(selected)}",
            flush=True,
        )

        for index, meta in enumerate(
            selected,
            1,
        ):
            if (
                index == 1
                or index % 10 == 0
                or index == len(selected)
            ):
                print(
                    "[CKPT0D] "
                    f"{source} "
                    f"table={index}/"
                    f"{len(selected)} "
                    f"path={meta.relpath}",
                    flush=True,
                )

            (
                summary,
                private,
            ) = analyze_table(
                meta
            )

            full_summaries.append(
                summary
            )

            private_by_key[
                (
                    source,
                    meta.relpath,
                )
            ] = private

    write_tsv(
        out
        / "full_table_audit.tsv",
        [
            "source",
            "path",
            "roles",
            "size_bytes",
            "row_count",
            "patient_id_column",
            "sample_id_column",
            "unique_patients",
            "unique_samples",
            "site_columns",
            "cancer_columns",
            "panel_columns",
            "date_columns",
            "parse_error",
        ],
        (
            [
                summary[
                    "source"
                ],
                summary[
                    "path"
                ],
                ",".join(
                    summary[
                        "roles"
                    ]
                ),
                summary[
                    "size_bytes"
                ],
                summary[
                    "row_count"
                ],
                (
                    summary[
                        "patient_id_column"
                    ]
                    or ""
                ),
                (
                    summary[
                        "sample_id_column"
                    ]
                    or ""
                ),
                summary[
                    "unique_patients"
                ],
                summary[
                    "unique_samples"
                ],
                ",".join(
                    summary[
                        "site_columns"
                    ]
                ),
                ",".join(
                    summary[
                        "cancer_columns"
                    ]
                ),
                ",".join(
                    summary[
                        "panel_columns"
                    ]
                ),
                ",".join(
                    summary[
                        "date_columns"
                    ]
                ),
                (
                    summary[
                        "parse_error"
                    ]
                    or ""
                ),
            ]
            for summary
            in full_summaries
        ),
    )

    ###########################################################################
    # Phase D: authoritative patient/sample sets.
    ###########################################################################

    source_patients: dict[
        str,
        set[str],
    ] = {
        key: set()
        for key
        in roots
    }

    source_samples: dict[
        str,
        set[str],
    ] = {
        key: set()
        for key
        in roots
    }

    for source in (
        "GENIE",
        "CHORD",
        "BPC",
    ):
        for meta in authority_metas[
            source
        ][
            "patient"
        ]:
            private = (
                private_by_key.get(
                    (
                        source,
                        meta.relpath,
                    )
                )
            )

            if private:
                source_patients[
                    source
                ].update(
                    private[
                        "patient_ids"
                    ]
                )

        for meta in authority_metas[
            source
        ][
            "sample"
        ]:
            private = (
                private_by_key.get(
                    (
                        source,
                        meta.relpath,
                    )
                )
            )

            if private:
                source_patients[
                    source
                ].update(
                    private[
                        "patient_ids"
                    ]
                )

                source_samples[
                    source
                ].update(
                    private[
                        "sample_ids"
                    ]
                )

    ###########################################################################
    # Phase E: CHORD breast cohort identification.
    ###########################################################################

    chord_breast: set[str] = set()

    chord_cancer_values: Counter[
        str
    ] = Counter()

    chord_evidence_files: list[
        str
    ] = []

    for summary in full_summaries:
        if (
            summary["source"]
            != "CHORD"
        ):
            continue

        if not summary[
            "cancer_columns"
        ]:
            continue

        private = private_by_key[
            (
                "CHORD",
                summary["path"],
            )
        ]

        if private[
            "breast_patients"
        ]:
            chord_breast.update(
                private[
                    "breast_patients"
                ]
            )

            chord_cancer_values.update(
                private[
                    "cancer_values"
                ]
            )

            chord_evidence_files.append(
                summary["path"]
            )

    ###########################################################################
    # Phase F: BPC MSK/DFCI/VICC partition.
    #
    # Explicit site columns are preferred.
    # Patient-ID prefix/site encoding is fallback only.
    #
    # Outcome distributions are never accumulated.
    ###########################################################################

    bpc_votes: dict[
        str,
        Counter[str],
    ] = defaultdict(Counter)

    bpc_raw_site_values: Counter[
        str
    ] = Counter()

    for summary in full_summaries:
        if (
            summary["source"]
            != "BPC"
        ):
            continue

        private = private_by_key[
            (
                "BPC",
                summary["path"],
            )
        ]

        for (
            pid,
            votes,
        ) in private[
            "patient_site_votes"
        ].items():
            bpc_votes[
                pid
            ].update(
                votes
            )

        bpc_raw_site_values.update(
            private[
                "site_values"
            ]
        )

    bpc_site_sets = {
        "MSK": set(),
        "DFCI": set(),
        "VICC": set(),
        "UNRESOLVED": set(),
        "CONFLICT": set(),
    }

    for pid in source_patients[
        "BPC"
    ]:
        votes = bpc_votes.get(
            pid,
            Counter(),
        )

        if not votes:
            inferred = (
                infer_site_from_patient_id(
                    pid
                )
            )

            if inferred:
                votes[
                    inferred
                ] += 1

        if not votes:
            bpc_site_sets[
                "UNRESOLVED"
            ].add(
                pid
            )

            continue

        ranked = votes.most_common()

        max_votes = (
            ranked[0][1]
        )

        winners = sorted(
            [
                site
                for site, count
                in ranked
                if count == max_votes
            ]
        )

        if len(winners) != 1:
            bpc_site_sets[
                "CONFLICT"
            ].add(
                pid
            )

        else:
            bpc_site_sets[
                winners[0]
            ].add(
                pid
            )

    bpc_partition = {
        "patient_counts": {
            key: len(value)
            for key, value
            in bpc_site_sets.items()
        },
        "raw_site_labels":
            bpc_raw_site_values
            .most_common(50),
        "policy_guard": (
            "DFCI/VICC are partitioned by identifiers/"
            "site labels only; no outcome-value "
            "distributions are computed or reported."
        ),
    }

    atomic_json(
        out
        / "bpc_site_partition.json",
        bpc_partition,
    )

    ###########################################################################
    # Phase G: exact cross-dataset patient-ID overlap.
    #
    # Exact means only whitespace trimming from input parsing.
    # No fuzzy matching or prefix stripping.
    ###########################################################################

    overlap_sets = {
        "GENIE_ALL":
            source_patients[
                "GENIE"
            ],

        "CHORD_ALL":
            source_patients[
                "CHORD"
            ],

        "CHORD_BREAST":
            chord_breast,

        "BPC_ALL":
            source_patients[
                "BPC"
            ],

        "BPC_MSK":
            bpc_site_sets[
                "MSK"
            ],

        "BPC_DFCI":
            bpc_site_sets[
                "DFCI"
            ],

        "BPC_VICC":
            bpc_site_sets[
                "VICC"
            ],
    }

    requested_pairs = [
        (
            "GENIE_ALL",
            "CHORD_ALL",
        ),
        (
            "GENIE_ALL",
            "CHORD_BREAST",
        ),
        (
            "GENIE_ALL",
            "BPC_ALL",
        ),
        (
            "GENIE_ALL",
            "BPC_MSK",
        ),
        (
            "GENIE_ALL",
            "BPC_DFCI",
        ),
        (
            "GENIE_ALL",
            "BPC_VICC",
        ),
        (
            "CHORD_ALL",
            "BPC_ALL",
        ),
        (
            "CHORD_BREAST",
            "BPC_ALL",
        ),
        (
            "CHORD_BREAST",
            "BPC_MSK",
        ),
        (
            "CHORD_BREAST",
            "BPC_DFCI",
        ),
        (
            "CHORD_BREAST",
            "BPC_VICC",
        ),
    ]

    overlaps = {
        f"{left}__{right}":
            pair_overlap(
                overlap_sets[
                    left
                ],
                overlap_sets[
                    right
                ],
            )
        for left, right
        in requested_pairs
    }

    overlaps[
        "definition"
    ] = (
        "Exact overlap after trimming surrounding "
        "whitespace only; no prefix stripping or "
        "fuzzy matching."
    )

    atomic_json(
        out
        / "cross_dataset_patient_overlap.json",
        overlaps,
    )

    ###########################################################################
    # Phase H: CHORD longitudinal domain inventory.
    ###########################################################################

    chord_domain_rows: list[
        list[Any]
    ] = []

    for summary in full_summaries:
        if (
            summary["source"]
            != "CHORD"
        ):
            continue

        roles = set(
            summary[
                "roles"
            ]
        )

        if not (
            roles
            & {
                "treatment",
                "radiology",
                "lab",
                "timeline",
                "outcome",
            }
        ):
            continue

        chord_domain_rows.append(
            [
                summary[
                    "path"
                ],
                ",".join(
                    summary[
                        "roles"
                    ]
                ),
                summary[
                    "row_count"
                ],
                summary[
                    "unique_patients"
                ],
                (
                    summary[
                        "patient_id_column"
                    ]
                    or ""
                ),
                ",".join(
                    summary[
                        "date_columns"
                    ]
                ),
                (
                    summary[
                        "parse_error"
                    ]
                    or ""
                ),
            ]
        )

    write_tsv(
        out
        / "chord_domain_summary.tsv",
        [
            "path",
            "roles",
            "row_count",
            "unique_patients",
            "patient_id_column",
            "date_columns",
            "parse_error",
        ],
        chord_domain_rows,
    )

    ###########################################################################
    # Phase I: GENIE sequencing-panel / assay structure.
    ###########################################################################

    genie_panel_counts: Counter[
        str
    ] = Counter()

    genie_panel_evidence: list[
        dict[str, Any]
    ] = []

    for summary in full_summaries:
        if (
            summary["source"]
            != "GENIE"
        ):
            continue

        if not summary[
            "panel_columns"
        ]:
            continue

        private = private_by_key[
            (
                "GENIE",
                summary["path"],
            )
        ]

        if private[
            "panel_values"
        ]:
            genie_panel_counts.update(
                private[
                    "panel_values"
                ]
            )

            genie_panel_evidence.append(
                {
                    "path":
                        summary[
                            "path"
                        ],
                    "panel_columns":
                        summary[
                            "panel_columns"
                        ],
                }
            )

    panel_definitions: list[
        dict[str, Any]
    ] = []

    for meta in all_metas[
        "GENIE"
    ]:
        definition = (
            summarize_gene_panel_definition(
                meta
            )
        )

        if definition:
            panel_definitions.append(
                definition
            )

    genie_panel = {
        "distinct_panel_values":
            len(
                genie_panel_counts
            ),

        "panel_value_counts_top100":
            genie_panel_counts
            .most_common(100),

        "panel_mapping_evidence":
            genie_panel_evidence,

        "panel_definition_files":
            panel_definitions,

        "coverage_interpretation": (
            "A sample-level panel/assay field plus "
            "gene-panel definitions is sufficient to "
            "construct explicit covered/uncovered gene "
            "masks. Matching IDs may require a small "
            "mapping step if assay values and definition "
            "IDs use different naming conventions."
        ),
    }

    atomic_json(
        out
        / "genie_panel_structure.json",
        genie_panel,
    )

    ###########################################################################
    # Phase J: timestamp candidates + source descriptions.
    #
    # This is schema-level evidence, not yet our final availability-time rule.
    ###########################################################################

    timestamp_rows: list[
        list[Any]
    ] = []

    for source in (
        "GENIE",
        "CHORD",
        "BPC",
    ):
        for meta in all_metas[
            source
        ]:
            descriptions = (
                cbioportal_column_descriptions(
                    meta
                )
            )

            for col in date_cols(
                meta
            ):
                timestamp_rows.append(
                    [
                        source,
                        meta.relpath,
                        ",".join(
                            meta.role
                        ),
                        col,
                        timestamp_semantic(
                            col
                        ),
                        descriptions.get(
                            col,
                            "",
                        ),
                    ]
                )

    write_tsv(
        out
        / "timestamp_candidates.tsv",
        [
            "source",
            "path",
            "roles",
            "column",
            "semantic_class",
            "source_description_if_available",
        ],
        timestamp_rows,
    )

    ###########################################################################
    # Phase K: readiness criteria.
    ###########################################################################

    role_presence: Counter[
        str
    ] = Counter()

    for summary in full_summaries:
        if (
            summary["source"]
            == "CHORD"
            and not summary[
                "parse_error"
            ]
        ):
            for role in summary[
                "roles"
            ]:
                role_presence[
                    role
                ] += 1

    genie_ok = (
        len(
            source_patients[
                "GENIE"
            ]
        )
        > 0
        and len(
            source_samples[
                "GENIE"
            ]
        )
        > 0
        and (
            len(
                genie_panel_counts
            )
            > 0
            or any(
                item.get(
                    "gene_count"
                )
                for item
                in panel_definitions
            )
        )
    )

    chord_ok = (
        len(
            source_patients[
                "CHORD"
            ]
        )
        > 0
        and len(
            chord_breast
        )
        > 0
        and role_presence[
            "treatment"
        ]
        > 0
        and role_presence[
            "radiology"
        ]
        > 0
    )

    bpc_ok = (
        all(
            len(
                bpc_site_sets[
                    site
                ]
            )
            > 0
            for site
            in (
                "MSK",
                "DFCI",
                "VICC",
            )
        )
        and len(
            bpc_site_sets[
                "CONFLICT"
            ]
        )
        == 0
    )

    dates_ok = any(
        row[0] == "CHORD"
        and (
            "treatment"
            in row[2]
            or "radiology"
            in row[2]
        )
        for row
        in timestamp_rows
    )

    # Soft sanity checks only. They do not gate readiness because minor
    # differences may reflect release organization or patient-ID unions.
    reference_counts = {
        "GENIE_patients":
            242866,
        "GENIE_samples":
            289869,
        "CHORD_patients":
            24950,
        "CHORD_breast_patients":
            5368,
        "BPC_patients":
            1130,
    }

    observed_reference = {
        "GENIE_patients":
            len(
                source_patients[
                    "GENIE"
                ]
            ),

        "GENIE_samples":
            len(
                source_samples[
                    "GENIE"
                ]
            ),

        "CHORD_patients":
            len(
                source_patients[
                    "CHORD"
                ]
            ),

        "CHORD_breast_patients":
            len(
                chord_breast
            ),

        "BPC_patients":
            len(
                source_patients[
                    "BPC"
                ]
            ),
    }

    reference_comparison = {
        key: {
            "reference":
                reference_counts[
                    key
                ],
            "observed":
                observed_reference[
                    key
                ],
            "delta":
                observed_reference[
                    key
                ]
                - reference_counts[
                    key
                ],
        }
        for key
        in reference_counts
    }

    readiness = {
        "status": (
            "READY_FOR_CANONICAL_TIMELINE"
            if (
                genie_ok
                and chord_ok
                and bpc_ok
                and dates_ok
            )
            else "NEEDS_SCHEMA_RESOLUTION"
        ),

        "criteria": {
            (
                "GENIE_patient_sample_"
                "and_panel_structure"
            ):
                genie_ok,

            (
                "CHORD_patient_breast_"
                "treatment_radiology_structure"
            ):
                chord_ok,

            (
                "BPC_MSK_DFCI_VICC_partition_"
                "without_conflict"
            ):
                bpc_ok,

            (
                "CHORD_treatment_or_radiology_"
                "timestamp_candidates"
            ):
                dates_ok,
        },

        "source_counts": {
            "GENIE": {
                "patients":
                    len(
                        source_patients[
                            "GENIE"
                        ]
                    ),
                "samples":
                    len(
                        source_samples[
                            "GENIE"
                        ]
                    ),
            },

            "CHORD": {
                "patients":
                    len(
                        source_patients[
                            "CHORD"
                        ]
                    ),
                "samples":
                    len(
                        source_samples[
                            "CHORD"
                        ]
                    ),
                "breast_patients":
                    len(
                        chord_breast
                    ),
            },

            "BPC": {
                "patients":
                    len(
                        source_patients[
                            "BPC"
                        ]
                    ),
                "samples":
                    len(
                        source_samples[
                            "BPC"
                        ]
                    ),
                "sites": {
                    key:
                        len(value)
                    for key, value
                    in bpc_site_sets.items()
                },
            },
        },

        "chord_breast_evidence_files":
            sorted(
                set(
                    chord_evidence_files
                )
            ),

        "chord_breast_top_cancer_labels":
            chord_cancer_values
            .most_common(30),

        "soft_reference_count_comparison":
            reference_comparison,

        "external_validation_policy": {
            "BPC_MSK": (
                "development-time schema/"
                "endpoint/label audit only"
            ),
            "BPC_DFCI": (
                "untouched external validation; "
                "no outcome distributions inspected"
            ),
            "BPC_VICC": (
                "untouched external validation; "
                "no outcome distributions inspected"
            ),
        },
    }

    atomic_json(
        out
        / "readiness.json",
        readiness,
    )

    ###########################################################################
    # Phase L: human-readable report.
    ###########################################################################

    lines: list[str] = []

    lines.append(
        "# Checkpoint 0D - "
        "Authoritative Local Schema Audit"
    )
    lines.append("")
    lines.append(
        "Status: "
        f"**{readiness['status']}**"
    )
    lines.append("")

    lines.append(
        "## Authoritative "
        "patient/sample identity"
    )
    lines.append("")

    for source in (
        "GENIE",
        "CHORD",
        "BPC",
    ):
        lines.append(
            f"### {source}"
        )
        lines.append("")

        lines.append(
            "- patients: "
            f"`{readiness['source_counts'][source]['patients']}`"
        )

        lines.append(
            "- samples: "
            f"`{readiness['source_counts'][source]['samples']}`"
        )

        lines.append(
            "- patient ID field(s): "
            f"`{', '.join(authority[source]['patient_id_fields'])}`"
        )

        lines.append(
            "- sample ID field(s): "
            f"`{', '.join(authority[source]['sample_id_fields'])}`"
        )

        lines.append(
            "- patient files:"
        )

        for path in authority[
            source
        ][
            "patient_files"
        ][:30]:
            lines.append(
                f"  - `{path}`"
            )

        lines.append(
            "- sample files:"
        )

        for path in authority[
            source
        ][
            "sample_files"
        ][:30]:
            lines.append(
                f"  - `{path}`"
            )

        lines.append("")

    lines.append(
        "## MSK-CHORD breast cohort"
    )
    lines.append("")

    lines.append(
        "- breast patients identified: "
        f"`{len(chord_breast)}`"
    )

    lines.append(
        "- evidence files: "
        f"`{', '.join(sorted(set(chord_evidence_files)))}`"
    )

    lines.append(
        "- top cancer labels in "
        "the evidence tables:"
    )

    for (
        label,
        count,
    ) in chord_cancer_values.most_common(
        20
    ):
        lines.append(
            f"  - `{label}`: {count}"
        )

    lines.append("")

    lines.append(
        "## BPC institution partition"
    )
    lines.append("")

    for site in (
        "MSK",
        "DFCI",
        "VICC",
        "UNRESOLVED",
        "CONFLICT",
    ):
        lines.append(
            f"- {site}: "
            f"`{len(bpc_site_sets[site])}` "
            "patients"
        )

    lines.append(
        "- Policy guard: DFCI/VICC "
        "outcome-value distributions "
        "were not computed."
    )
    lines.append("")

    lines.append(
        "## Exact patient-ID overlaps"
    )
    lines.append("")

    lines.append(
        "Exact means whitespace-trimmed "
        "string equality only; no fuzzy "
        "matching or prefix stripping."
    )
    lines.append("")

    for key, value in overlaps.items():
        if key == "definition":
            continue

        lines.append(
            f"- {key}: "
            f"`{value['exact_overlap_count']}`"
        )

    lines.append("")

    lines.append(
        "## GENIE panel structure"
    )
    lines.append("")

    lines.append(
        "- distinct panel/assay "
        "values observed: "
        f"`{len(genie_panel_counts)}`"
    )

    lines.append(
        "- panel definition-like "
        "files inspected: "
        f"`{len(panel_definitions)}`"
    )

    lines.append(
        "- top panel/assay values:"
    )

    for (
        label,
        count,
    ) in genie_panel_counts.most_common(
        30
    ):
        lines.append(
            f"  - `{label}`: {count}"
        )

    lines.append("")

    lines.append(
        "## CHORD longitudinal "
        "domain tables"
    )
    lines.append("")

    for summary in full_summaries:
        if (
            summary["source"]
            != "CHORD"
        ):
            continue

        if not (
            set(
                summary[
                    "roles"
                ]
            )
            & {
                "treatment",
                "radiology",
                "lab",
                "timeline",
                "outcome",
            }
        ):
            continue

        lines.append(
            f"- `{summary['path']}` "
            f"| roles={','.join(summary['roles'])} "
            f"| rows={summary['row_count']} "
            f"| patients={summary['unique_patients']} "
            f"| dates={','.join(summary['date_columns'])}"
        )

    lines.append("")

    lines.append(
        "## Timestamp candidates"
    )
    lines.append("")

    lines.append(
        "- schema-level timestamp "
        "candidates: "
        f"`{len(timestamp_rows)}`"
    )

    lines.append(
        "- Full mapping: "
        "`timestamp_candidates.tsv`"
    )
    lines.append("")

    lines.append(
        "## Soft reference-count comparison"
    )
    lines.append("")

    lines.append(
        "These are sanity checks only, "
        "not readiness gates; discrepancies "
        "can reflect release structure or "
        "ID-union rules."
    )
    lines.append("")

    for (
        key,
        value,
    ) in reference_comparison.items():
        lines.append(
            f"- {key}: "
            f"observed=`{value['observed']}` "
            f"reference=`{value['reference']}` "
            f"delta=`{value['delta']}`"
        )

    lines.append("")

    lines.append(
        "## Readiness criteria"
    )
    lines.append("")

    for (
        key,
        value,
    ) in readiness[
        "criteria"
    ].items():
        lines.append(
            f"- {key}: "
            f"**{'PASS' if value else 'FAIL'}**"
        )

    lines.append("")

    lines.append(
        "## Artifacts"
    )
    lines.append("")

    for name in (
        "schema_inventory.jsonl",
        "authoritative_entities.json",
        "full_table_audit.tsv",
        "bpc_site_partition.json",
        "cross_dataset_patient_overlap.json",
        "chord_domain_summary.tsv",
        "genie_panel_structure.json",
        "timestamp_candidates.tsv",
        "readiness.json",
    ):
        lines.append(
            f"- `{name}`"
        )

    lines.append("")

    atomic_text(
        out
        / "audit.md",
        "\n".join(lines)
        + "\n",
    )

    ###########################################################################
    # Phase M: persistent handoff.
    ###########################################################################

    handoff = {
        "checkpoint": "0D",

        "name": (
            "authoritative_schema_overlap_"
            "and_readiness_audit"
        ),

        "status":
            readiness[
                "status"
            ],

        "source_counts":
            readiness[
                "source_counts"
            ],

        "readiness_criteria":
            readiness[
                "criteria"
            ],

        "authoritative_entities_artifact":
            (
                "artifacts/checkpoint0d/"
                "authoritative_entities.json"
            ),

        "cross_dataset_overlap_artifact":
            (
                "artifacts/checkpoint0d/"
                "cross_dataset_patient_overlap.json"
            ),

        "bpc_partition_artifact":
            (
                "artifacts/checkpoint0d/"
                "bpc_site_partition.json"
            ),

        "genie_panel_artifact":
            (
                "artifacts/checkpoint0d/"
                "genie_panel_structure.json"
            ),

        "chord_domain_artifact":
            (
                "artifacts/checkpoint0d/"
                "chord_domain_summary.tsv"
            ),

        "timestamp_artifact":
            (
                "artifacts/checkpoint0d/"
                "timestamp_candidates.tsv"
            ),

        "policy": {
            "GENIE":
                "genomic pretraining",

            "CHORD":
                (
                    "temporal pretraining + "
                    "primary dynamic training"
                ),

            "BPC_MSK":
                (
                    "schema/endpoint/label audit "
                    "only during development"
                ),

            "BPC_DFCI":
                (
                    "untouched external "
                    "validation"
                ),

            "BPC_VICC":
                (
                    "untouched external "
                    "validation"
                ),
        },

        "next_required_action": (
            "Begin canonical timeline construction"
            if (
                readiness[
                    "status"
                ]
                == "READY_FOR_CANONICAL_TIMELINE"
            )
            else (
                "Resolve failed readiness criteria "
                "before canonical timeline construction"
            )
        ),
    }

    handoff_dir = (
        repo
        / "artifacts"
        / "handoff"
    )

    handoff_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    atomic_json(
        handoff_dir
        / "checkpoint_00D.json",
        handoff,
    )

    ###########################################################################
    # Concise terminal output for the next agent/chat turn.
    ###########################################################################

    print("")
    print(
        "========== CKPT0D SUMMARY =========="
    )

    print(
        "status="
        f"{readiness['status']}"
    )

    print(
        "GENIE "
        f"patients={len(source_patients['GENIE'])} "
        f"samples={len(source_samples['GENIE'])} "
        f"panels={len(genie_panel_counts)}"
    )

    print(
        "CHORD "
        f"patients={len(source_patients['CHORD'])} "
        f"samples={len(source_samples['CHORD'])} "
        f"breast_patients={len(chord_breast)}"
    )

    print(
        "BPC "
        + " ".join(
            f"{key}={len(bpc_site_sets[key])}"
            for key in (
                "MSK",
                "DFCI",
                "VICC",
                "UNRESOLVED",
                "CONFLICT",
            )
        )
    )

    for (
        key,
        value,
    ) in readiness[
        "criteria"
    ].items():
        print(
            "criterion "
            f"{key}="
            f"{'PASS' if value else 'FAIL'}"
        )

    print(
        "overlaps:"
    )

    for (
        key,
        value,
    ) in overlaps.items():
        if key == "definition":
            continue

        print(
            f"  {key}="
            f"{value['exact_overlap_count']}"
        )

    print(
        "chord_domain_tables="
        f"{len(chord_domain_rows)}"
    )

    print(
        "timestamp_candidates="
        f"{len(timestamp_rows)}"
    )

    print(
        "policy_guard_external_"
        "outcome_distributions=PASS"
    )

    print(
        "soft_reference_counts:"
    )

    for (
        key,
        value,
    ) in reference_comparison.items():
        print(
            f"  {key}: "
            f"observed={value['observed']} "
            f"reference={value['reference']} "
            f"delta={value['delta']}"
        )

    print(
        "audit_md="
        "artifacts/checkpoint0d/audit.md"
    )

    print(
        "handoff="
        "artifacts/handoff/checkpoint_00D.json"
    )

    print(
        "========== CKPT0D SUMMARY END =========="
    )

    print("")
    print(
        "========== CKPT0D DECISION PACKET =========="
    )

    print(
        "AUTHORITATIVE_IDENTITY:"
    )

    for source in (
        "GENIE",
        "CHORD",
        "BPC",
    ):
        print(
            f"  {source} "
            "patient_id_fields="
            f"{authority[source]['patient_id_fields']} "
            "sample_id_fields="
            f"{authority[source]['sample_id_fields']}"
        )

        for path in authority[
            source
        ][
            "patient_files"
        ][:12]:
            print(
                "    patient_file="
                f"{path}"
            )

        for path in authority[
            source
        ][
            "sample_files"
        ][:12]:
            print(
                "    sample_file="
                f"{path}"
            )

    print(
        "BPC_PARTITION:"
    )

    for site in (
        "MSK",
        "DFCI",
        "VICC",
        "UNRESOLVED",
        "CONFLICT",
    ):
        print(
            f"  {site}="
            f"{len(bpc_site_sets[site])}"
        )

    for (
        label,
        count,
    ) in bpc_raw_site_values.most_common(
        20
    ):
        print(
            "  raw_site_label="
            f"{label!r} "
            f"rows={count}"
        )

    print(
        "CHORD_BREAST:"
    )

    print(
        "  patients="
        f"{len(chord_breast)}"
    )

    for path in sorted(
        set(
            chord_evidence_files
        )
    ):
        print(
            "  evidence_file="
            f"{path}"
        )

    for (
        label,
        count,
    ) in chord_cancer_values.most_common(
        20
    ):
        print(
            "  cancer_label="
            f"{label!r} "
            f"rows={count}"
        )

    print(
        "CHORD_DOMAIN_TABLES:"
    )

    for row in chord_domain_rows:
        print(
            "  "
            + " | ".join(
                str(x)
                for x in row
            )
        )

    print(
        "GENIE_PANEL_STRUCTURE:"
    )

    print(
        "  distinct_panel_values="
        f"{len(genie_panel_counts)}"
    )

    print(
        "  panel_definition_files="
        f"{len(panel_definitions)}"
    )

    for (
        label,
        count,
    ) in genie_panel_counts.most_common(
        30
    ):
        print(
            "  panel="
            f"{label!r} "
            f"rows={count}"
        )

    for item in panel_definitions[
        :30
    ]:
        print(
            "  panel_definition="
            f"{item.get('path')} "
            "panel_id="
            f"{item.get('panel_id')} "
            "gene_count="
            f"{item.get('gene_count')}"
        )

    print(
        "CHORD_TIMESTAMPS:"
    )

    shown = 0

    for row in timestamp_rows:
        (
            source,
            path,
            roles,
            col,
            semantic,
            description,
        ) = row

        if source != "CHORD":
            continue

        if not any(
            role in roles
            for role in (
                "treatment",
                "radiology",
                "lab",
                "timeline",
                "outcome",
                "patient",
                "sample",
            )
        ):
            continue

        print(
            f"  {path} "
            f"| {roles} "
            f"| {col} "
            f"| {semantic} "
            f"| {description}"
        )

        shown += 1

        if shown >= 100:
            print(
                "  ... additional "
                "timestamp candidates are "
                "in timestamp_candidates.tsv"
            )
            break

    print(
        "EXACT_OVERLAPS:"
    )

    for (
        key,
        value,
    ) in overlaps.items():
        if key == "definition":
            continue

        print(
            f"  {key}="
            f"{value['exact_overlap_count']}"
        )

    print(
        "========== CKPT0D DECISION PACKET END =========="
    )

    return 0


if __name__ == "__main__":
    raise SystemExit(
        main()
    )
